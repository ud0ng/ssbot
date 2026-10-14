"""구글 시트 템플릿(.xlsx) 만들기.

  python tools/make_templates.py      → templates/*.xlsx

구글 드라이브에 올리면 구글 시트로 열린다. 내용은 tests/fixtures 의 예시 시트를 바탕으로 한다.
"""

from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
FIXTURES = ROOT / 'tests' / 'fixtures'
OUT = ROOT / 'templates'

HEADER_FILL = PatternFill('solid', fgColor='CFE2F3')
NOTE_FILL = PatternFill('solid', fgColor='FFF2CC')
HEADER_FONT = Font(bold=True)


def read(sheet: str, tab: str) -> list:
    with (FIXTURES / sheet / f'{tab}.csv').open(encoding='utf-8') as file:
        return [row for row in csv.reader(file)]


def add(workbook: Workbook, title: str, rows: list, header_row: int = 1, widths: dict | None = None, notes: tuple = ()) -> None:
    sheet = workbook.create_sheet(title)
    for row in rows:
        sheet.append(row)
    for cell in sheet[header_row]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
    for index in notes:
        for cell in sheet[index]:
            cell.fill = NOTE_FILL
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical='top')
    for col in range(1, sheet.max_column + 1):
        letter = get_column_letter(col)
        longest = max((len(str(c.value or '').split('\n')[0]) for c in sheet[letter]), default=8)
        sheet.column_dimensions[letter].width = (widths or {}).get(letter, min(max(10, longest * 1.6), 70))
    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1)


def new_workbook() -> Workbook:
    workbook = Workbook()
    workbook.remove(workbook.active)
    return workbook


def main_sheet() -> Workbook:
    workbook = new_workbook()

    help_rows = [['명령어', '설명']] + [[c, d] for c, d in [
        ('[도움말]', '도움말을 보여줍니다.'),
        ('[NdM]', 'M면체 주사위를 N개 굴립니다. 예: [1d100]'),
        ('[랜덤/A, B, C...]', '여러 옵션 중 1개를 골라 출력합니다.'),
        ('[운세]', '오늘의 운세를 확인합니다.'),
        ('[YN]', ''),
        ('[가위바위보]', ''),
        ('[홀짝]', ''),
        ('[현자의 조언]', ''),
        # 가방·상점·구매·양도·출석 같은 재화 명령어 도움말은 마스토돈(컨텐츠 설정 > 상점)에서 붙는다
        ('[랜덤박스/박스명]', '랜덤박스를 엽니다. [랜덤박스]는 박스 목록'),
        ('[장소 목록]', '진입할 수 있는 장소 목록을 봅니다.'),
        ('[진입/장소명]', '장소에 들어갑니다.'),
        ('[조사/포인트명]', '지금 장소의 조사 포인트를 조사합니다.'),
        ('[쪽지시험]', '오늘의 쪽지시험 점수를 확인합니다.'),
        ('[낚시]', '강가에서 낚시합니다.'),
        ('[스킬]', '스킬을 사용합니다.'),
    ]]
    add(workbook, '도움말', help_rows, widths={'A': 26, 'B': 60})

    admin_rows = [['명령어', '설명', '예시'], *[[a, b, c] for a, b, c in [
        ('[캐시 리셋]', '시트를 고친 내용을 바로 반영합니다. (안 해도 1분 안에 반영)', '[캐시 리셋]'),
        # 소지금·아이템·스탯 변경 명령어 도움말은 마스토돈(컨텐츠 설정 > 상점)에 있다
        ('[스진/워크시트명]', 'STORY 계정에 DM 으로 보내면 스토리 시트의 그 워크시트를 송출', '[스진/1일차]'),
        ('[스진 중지]', '진행 중인 스토리를 멈춥니다.', '[스진 중지]'),
    ]]]
    add(workbook, '운영진 명령어', admin_rows, widths={'A': 34, 'B': 60, 'C': 34})

    # 명단·관리·상점·출석부·기록은 마스토돈(캐릭터 관리, 컨텐츠 설정 > 상점)으로 옮겼다
    add(workbook, '운세', read('main', '운세'), widths={'A': 100})

    keywords = read('main', '키워드')
    keywords[0] = ['키워드', '답변', '이미지']
    keywords.insert(1, ['고양이', '야옹! {시전자}{을를} 쳐다본다.', '(이미지 주소를 넣으면 답변에 첨부됩니다. 예: https://…/cat.png)'])
    add(workbook, '키워드', keywords, widths={'A': 18, 'B': 70, 'C': 40})

    add(workbook, '랜덤박스', read('main', '랜덤박스'), widths={'A': 14, 'B': 8, 'C': 16, 'D': 8, 'E': 8})
    add(workbook, '설정', read('main', '설정'), widths={'A': 18, 'B': 16, 'C': 60})
    workbook.move_sheet('설정', offset=-workbook.sheetnames.index('설정'))
    return workbook


def reservation_sheet() -> Workbook:
    workbook = new_workbook()
    rows = read('reservation', '관리')
    add(workbook, '관리', rows, widths={'A': 10, 'B': 8, 'C': 10, 'D': 8, 'E': 80})
    return workbook


def story_sheet() -> Workbook:
    workbook = new_workbook()
    add(workbook, '테스트', read('story', '테스트'), header_row=2, notes=(1,), widths={'A': 12, 'B': 8, 'C': 80})
    add(workbook, '예시 파트1', read('story', '예시 파트1'), widths={'A': 12, 'B': 8, 'C': 80})
    add(workbook, '빈칸', [['계정', '간격', '문구']], widths={'A': 12, 'B': 8, 'C': 80})
    return workbook


def investigation_sheet() -> Workbook:
    workbook = new_workbook()
    entry = read('investigation', '진입')
    add(workbook, '진입', entry, notes=(2, 3, 4), widths={'A': 14, 'B': 90, 'C': 14})
    for place in ('운동장', '교실', '음악실'):
        add(workbook, place, read('investigation', place), widths={'A': 14, 'B': 80, 'C': 22, 'D': 22})
    return workbook


def main() -> None:
    OUT.mkdir(exist_ok=True)
    for name, build in (('1_기본_상점_스탯', main_sheet), ('2_예약툿', reservation_sheet),
                        ('3_스토리', story_sheet), ('4_조사', investigation_sheet)):
        path = OUT / f'{name}.xlsx'
        build().save(path)
        print('saved', path.name)


if __name__ == '__main__':
    main()
